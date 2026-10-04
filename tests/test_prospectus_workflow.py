"""Prospectus upload -> read -> check -> publish (reallocate) -> undo.

Every publish here replaces the whole live curriculum, so each test that
publishes undoes it again before finishing and leaves the seed data intact.
All curriculum values are DEMO data.
"""

import io
import itertools
from decimal import Decimal

import pytest
from flask import has_app_context

import curriculum_service as svc
from conftest import db
from models import (
    CurriculumEntry, CurriculumVersion, Department, DraftProgramme, LecturerAssignment, Module,
    ProspectusChunk, ProspectusDocument, Resource, Topic, User,
)

_ids = itertools.count(1)
HEADER = "department,programme,nta_level,year_label,semester,module_code,module_name,module_type,credits,prerequisites,description\n"


def _unique(prefix="DEMO"):
    return f"{prefix} {next(_ids):03d}"


def _csv(rows):
    body = "".join(f"{d},{p},{level},,{sem},{code},{name},core,{credits},{prereq},DEMO row\n"
                   for d, p, level, sem, code, name, credits, prereq in rows)
    return (HEADER + body).encode("utf-8")


def _upload(admin, content, filename="demo.csv", *, year="2031/2032", large=True):
    title = f"{_unique()} prospectus"
    data = {"title": title, "academic_year_label": year, "file": (io.BytesIO(content), filename)}
    if large:
        data["allow_large_change"] = "1"
    response = admin.post("/admin/prospectus", data=data, content_type="multipart/form-data")
    assert response.status_code == 302
    return title


def _version(app, title):
    if has_app_context():
        return CurriculumVersion.query.filter(CurriculumVersion.label.startswith(title)).one()
    with app.app_context():
        version = CurriculumVersion.query.filter(CurriculumVersion.label.startswith(title)).one()
        return version


def _live_snapshot(app):
    with app.app_context():
        modules = {m.id: (m.publication_status, m.is_active) for m in Module.query.all()}
        resources = {r.id: r.module_id for r in Resource.query.all()}
        departments = {d.id: d.is_active for d in Department.query.all()}
        students = {u.id: (u.programme_id, u.nta_level_id, u.semester_id, u.academic_year_id)
                    for u in User.query.filter_by(role="student")}
        return modules, resources, departments, students


@pytest.fixture
def live_restored(app, admin):
    """Undo every publish a test makes and prove the seed curriculum is back."""
    before = _live_snapshot(app)
    with app.app_context():
        live_before = getattr(svc.live_version(), "id", None)
    yield
    with app.app_context():
        while getattr(svc.live_version(), "id", None) != live_before:
            svc.undo_last_publish(user=None)
            db.session.commit()
    after = _live_snapshot(app)
    for old, new in zip(before, after):
        assert {key: new[key] for key in old} == old


@pytest.fixture
def seeded_paths(app, seeded):
    """Give the seeded teaching module a DEMO code so the prospectus can match it."""
    with app.app_context():
        module = db.session.get(Module, seeded["teaching_module_id"])
        module.code = f"DEMO-T{next(_ids):03d}"
        student = User.query.filter_by(email="student@dit.ac.tz").one()
        db.session.commit()
        paths = {
            "code": module.code, "module_id": module.id,
            "department": student.department.name, "programme": student.programme.name,
            "level": student.nta_level.level_number, "semester": student.semester.semester_number,
            "student_id": student.id,
        }
    yield paths
    with app.app_context():
        db.session.get(Module, paths["module_id"]).code = None
        db.session.commit()


def test_clean_upload_goes_live_moves_resources_and_can_be_undone(app, admin, seeded_paths, live_restored):
    p = seeded_paths
    with app.app_context():
        resources = [r.id for r in Resource.query.filter_by(module_id=p["module_id"])]
        assignment = LecturerAssignment.query.filter_by(module_id=p["module_id"], status="approved").first()
        assignment_id = assignment.id
        topics = [t.id for t in Topic.query.filter_by(module_id=p["module_id"])]
        live_before = Module.query.filter(Module.publication_status != "archived").count()
    title = _upload(admin, _csv([
        ("DEMO Moved Dept", "DEMO Programme", 4, 1, p["code"], "DEMO Taught Module", 10, ""),
        (p["department"], p["programme"], p["level"], p["semester"], "DEMO-NEW1", "DEMO New Module", 8, ""),
        (p["department"], p["programme"], p["level"], p["semester"], "DEMO-NEW2", "DEMO Next Module", "", "DEMO-NEW1"),
    ]))
    with app.app_context():
        version = _version(app, title)
        assert version.status == "published", version.validation_errors
        report = version.allocation_report
        assert report["counts"]["moved"] == 1 and report["counts"]["new"] == 2
        assert report["counts"]["retired"] == live_before - 1
        assert "DEMO Moved Dept" not in report["hidden_departments"]

        old = db.session.get(Module, p["module_id"])
        assert old.publication_status == "archived" and not old.is_active
        new = Module.query.filter_by(code=p["code"], curriculum_version_id=version.id).one()
        assert new.is_published and new.provenance == "prospectus"
        assert new.semester.nta_level.programme.department.name == "DEMO Moved Dept"
        assert {r.module_id for r in Resource.query.filter(Resource.id.in_(resources))} == {new.id}
        moved = db.session.get(LecturerAssignment, assignment_id)
        assert moved.module_id == new.id and moved.status == "approved"
        assert all(db.session.get(Topic, t).module_id == new.id for t in topics)
        nxt = Module.query.filter_by(code="DEMO-NEW2", curriculum_version_id=version.id).one()
        assert [link.prerequisite.code for link in nxt.prerequisite_links] == ["DEMO-NEW1"]

        student = db.session.get(User, p["student_id"])
        assert student.academic_year.label == "2031/2032" and student.academic_year.is_current
        assert Module.query.filter_by(semester_id=student.semester_id, is_active=True).count() == 2

    assert admin.post("/admin/prospectus/undo").status_code == 302
    with app.app_context():
        assert db.session.get(CurriculumVersion, version.id).status == "undone"
        assert db.session.get(Module, p["module_id"]).is_published
        assert {r.module_id for r in Resource.query.filter(Resource.id.in_(resources))} == {p["module_id"]}
        assert db.session.get(LecturerAssignment, assignment_id).module_id == p["module_id"]
        assert not Department.query.filter_by(name="DEMO Moved Dept").one().is_active
        assert db.session.get(User, p["student_id"]).academic_year.label != "2031/2032"


def test_result_page_shows_the_report(app, admin, hod, seeded_paths, live_restored):
    p = seeded_paths
    title = _upload(admin, _csv([(p["department"], p["programme"], p["level"], p["semester"], p["code"],
                                "DEMO Kept Module", 10, "")]))
    version = _version(app, title)
    page = admin.get(f"/admin/prospectus/{version.id}")
    assert page.status_code == 200 and b"Carried over" in page.data
    assert b"What went live" in page.data and p["programme"].encode() in page.data
    assert hod.get("/department/curriculum").data.count(b"DEMO Kept Module") == 1


def test_shared_code_without_a_programme_match_needs_placing(app, admin, seeded_paths, live_restored):
    p = seeded_paths
    title = _upload(admin, _csv([
        ("DEMO Dept A", "DEMO Programme A", 4, 1, p["code"], "DEMO Shared Module", 10, ""),
        ("DEMO Dept B", "DEMO Programme B", 4, 1, p["code"], "DEMO Shared Module", 10, ""),
    ]))
    with app.app_context():
        version = _version(app, title)
        assert version.allocation_report["counts"]["unplaced"] == 1
        assert Resource.query.filter_by(module_id=p["module_id"]).count() >= 1
        lecturer_modules = {a.module_id for a in LecturerAssignment.query.filter(
            LecturerAssignment.module.has(code=p["code"]), LecturerAssignment.status == "approved")}
        new_ids = {m.id for m in Module.query.filter_by(code=p["code"], curriculum_version_id=version.id)}
        assert new_ids <= lecturer_modules and p["module_id"] in lecturer_modules


@pytest.mark.parametrize("content,filename,message", [
    (b"%PDF-1.4\n%%EOF", "scanned.pdf", "No readable text"),
    (HEADER.encode() + b"DEMO Dept,DEMO Prog,4,,1,,DEMO Nameless Code,core,,,\n", "demo.csv", "module code is required"),
    (HEADER.encode() + b"DEMO Dept,DEMO Prog,4,,1,DEMO-P1,DEMO Loop,core,,DEMO-P9,\n", "demo.csv", "not in this prospectus"),
])
def test_a_bad_read_publishes_nothing(app, admin, content, filename, message):
    before = _live_snapshot(app)
    title = _upload(admin, content, filename)
    with app.app_context():
        version = _version(app, title)
        assert version.status == "failed"
        assert message in " ".join(issue["message"] for issue in version.validation_errors)
    assert _live_snapshot(app) == before
    assert message.encode() in admin.get(f"/admin/prospectus/{version.id}").data


def test_large_change_needs_confirmation(app, admin):
    before = _live_snapshot(app)
    title = _upload(admin, _csv([("DEMO Dept", "DEMO Prog", 4, 1, "DEMO-L1", "DEMO Lone Module", 5, "")]), large=False)
    with app.app_context():
        errors = _version(app, title).validation_errors
        assert any("would retire" in issue["message"] for issue in errors)
    assert _live_snapshot(app) == before


def test_same_file_twice_is_refused(app, admin, live_restored):
    content = _csv([("DEMO Dept", "DEMO Prog", 4, 1, "DEMO-S1", "DEMO Same Module", 5, "")])
    assert _version(app, _upload(admin, content)).status == "published"
    second = _version(app, _upload(admin, content))
    assert second.status == "failed" and "already live" in second.validation_errors[0]["message"]


def test_undo_is_refused_once_lecturers_add_content(app, admin, seeded_paths, live_restored):
    p = seeded_paths
    title = _upload(admin, _csv([(p["department"], p["programme"], p["level"], p["semester"], p["code"],
                                "DEMO Busy Module", 10, "")]))
    with app.app_context():
        version = _version(app, title)
        module = Module.query.filter_by(code=p["code"], curriculum_version_id=version.id).one()
        topic = Topic(module_id=module.id, title="DEMO topic added after publish")
        db.session.add(topic)
        db.session.commit()
        with pytest.raises(svc.CurriculumWorkflowError, match="can no longer be undone"):
            svc.undo_last_publish(user=None)
        db.session.rollback()
        db.session.delete(db.session.get(Topic, topic.id))
        db.session.commit()


def test_upload_keeps_the_original_file(app, admin):
    title = _upload(admin, b"%PDF-1.4\n%%EOF", "kept.pdf")
    with app.app_context():
        document = ProspectusDocument.query.filter_by(title=title).one()
        assert document.sha256 and document.extraction_status == "failed"
        version_id = _version(app, title).id
    download = admin.get(f"/admin/prospectus/{version_id}/file")
    assert download.status_code == 200 and download.data.startswith(b"%PDF")


def test_invalid_file_type_is_rejected(app, admin):
    response = admin.post("/admin/prospectus", data={
        "title": "DEMO bad type", "academic_year_label": "2031/2032",
        "file": (io.BytesIO(b"MZ\x90\x00"), "prospectus.exe"),
    }, content_type="multipart/form-data", follow_redirects=True)
    assert b"PDF, DOCX, TXT, MD or CSV" in response.data


def test_text_extraction_reads_structure(app, admin):
    text = (
        "DEPARTMENT OF DEMO STUDIES\n"
        "Ordinary Diploma in DEMO Technology\n"
        "NTA Level 6 Semester I\n"
        "DMO 06101 DEMO Circuit Theory 12\n"
        "DMO 06102 - DEMO Workshop Practice\n"
        "Semester II\n"
        "DMO 06201 DEMO Measurements 9.5\n"
    )
    candidates = svc.extract_candidates(text)
    assert [c["module_code"] for c in candidates] == ["DMO 06101", "DMO 06102", "DMO 06201"]
    assert candidates[0]["department"] == "Demo Studies"
    assert candidates[0]["programme"] == "DEMO Technology"
    assert candidates[0]["year_label"] == "Ordinary Diploma"
    assert (candidates[0]["nta_level"], candidates[0]["semester"]) == (6, 1)
    assert candidates[2]["semester"] == 2 and candidates[2]["credits"] == "9.5"
    assert candidates[1]["credits"] is None

    title = _upload(admin, text.encode(), "prospectus.txt", large=False)
    with app.app_context():
        version = _version(app, title)
        assert len(version.entries) == 3 and {e.origin for e in version.entries} == {"extracted"}
        assert ProspectusChunk.query.filter_by(document_id=version.prospectus_document_id).count() >= 1


# DEMO text laid out the way the DIT prospectus prints its programme tables.
DIT_LAYOUT = """\
6.9 DEPARTMENT OF DEMO ENGINEERING
6.9.1 Programmes offered by the Department of DEMO Engineering
(a) BASIC TECHNICIAN CERTIFICATE (BTC) IN DEMO ENGINEERING (NTA
LEVEL 4)
Semester I
Module
Code Module Title Credit
FUNDAMENTAL MODULE
GST 04111 DEMO Algebra 6
DIT Prospectus Academic Year 2025/2026
76 | P a g e
CORE MODULES
DMO04112 DEMO Workshop Practice and
Safety
9
GST 04111 DEMO Algebra 6
Total 15
SEMISTER II
DMO 04211 DEMO Surveying 12
(b). HIGHER DIPLOMA IN DEMO ENGINEERING - NTA LEVEL 7 [OLD
Semester IV Modules
S/N Module Name Class Credits
1. DMU 07411 DEMO Structures Core 9
2. DMU 07412 DEMO Communication Fundamental 6
(c). GENERAL COURSE PROGRAMME IN DEMO ENGINEERING
Semester I
DMG 4101 DEMO Bridging Module 9
8.9 DEMO CAMPUS
(a) ORDINARY DIPLOMA IN DEMO ENGINEERING (NTA 6)
SEMESTER I
SLT P 06101 DEMO Electromagnetism 2
6.9.2 List of Academic Staff in the Department of DEMO Engineering
TZS 15000 DEMO fee line 5
"""


def test_reader_follows_the_dit_prospectus_layout():
    rows, notes = svc.extract_prospectus(DIT_LAYOUT)
    by_code = {row["module_code"]: row for row in rows}
    assert list(by_code) == ["GST 04111", "DMO 04112", "DMO 04211", "DMU 07411", "DMU 07412", "SLTP 06101"]
    first = by_code["GST 04111"]
    assert (first["department"], first["programme"], first["nta_level"], first["semester"]) == \
        ("Demo Engineering", "Demo Engineering", 4, 1)
    assert first["year_label"] == "Basic Technician Certificate" and first["module_type"] == "fundamental"
    workshop = by_code["DMO 04112"]
    assert workshop["module_name"] == "DEMO Workshop Practice and Safety" and workshop["credits"] == "9"
    assert workshop["module_type"] == "core"
    assert by_code["DMO 04211"]["semester"] == 2
    structures, communication = by_code["DMU 07411"], by_code["DMU 07412"]
    assert (structures["nta_level"], structures["semester"], structures["module_name"]) == (7, 4, "DEMO Structures")
    assert structures["year_label"] == "Higher Diploma"
    assert communication["module_type"] == "fundamental"
    campus = by_code["SLTP 06101"]
    assert (campus["department"], campus["programme"], campus["nta_level"]) == ("Demo Campus", "Demo Engineering", 6)
    assert len(notes) == 1 and "General Course Programme" in notes[0].title()


def test_whole_prospectus_is_read_not_just_its_opening(app, admin, live_restored):
    padding = "DEMO introduction text that fills the opening chapters.\n" * 3000
    title = _upload(admin, (padding + DIT_LAYOUT).encode(), "prospectus.txt", large=True)
    with app.app_context():
        version = _version(app, title)
        assert version.status == "published", version.validation_errors
        assert len(version.entries) == 6
        assert any("no NTA level" in w["message"] for w in version.validation_warnings)


def test_prospectus_passages_carry_their_heading():
    text = ("CHAPTER FIVE\nEXAMINATION REGULATIONS\n9.0 Absence from Examination\n"
            "9.1 A DEMO candidate who absents oneself from a scheduled examination is discontinued.\n"
            "DIT Prospectus Academic Year 2025/2026\n52 | P a g e\n10.0 Postponement of Examination\n"
            "10.1 DEMO postponement needs approval from the Head of Department.\n")
    chunks = svc.prospectus_chunks(text)
    assert chunks[0].startswith("[Chapter Five: Examination Regulations › 9.0 Absence from Examination]")
    assert chunks[1].startswith("[Chapter Five: Examination Regulations › 10.0 Postponement of Examination]")
    assert not any("P a g e" in chunk for chunk in chunks)


def test_csv_missing_columns_is_rejected():
    with pytest.raises(svc.CurriculumWorkflowError, match="missing required columns"):
        svc.parse_curriculum_csv("department,programme\nA,B\n")


# --- validation rules --------------------------------------------------------

@pytest.fixture
def draft(app):
    """A staged version with one programme (levels 4-5)."""
    with app.app_context():
        version = CurriculumVersion(label=_unique(), academic_year_label="2033/2034", status="failed")
        programme = DraftProgramme(version=version, department_name=f"{version.label} Dept",
                                   name=f"{version.label} Programme", levels_csv="4,5", semesters_per_level=2)
        db.session.add_all([version, programme])
        db.session.flush()
        yield version, programme
        db.session.rollback()


def _entry(version, programme, code, name="DEMO Module", level=4, semester=1, **extra):
    entry = CurriculumEntry(version=version, programme=programme, module_code=code, module_name=name,
                            nta_level=level, semester_number=semester,
                            credits=extra.pop("credits", Decimal("10")), **extra)
    db.session.add(entry)
    db.session.flush()
    return entry


def _messages(report):
    return " | ".join(issue["message"] for issue in report.errors)


def test_printed_duplicate_codes_are_notes_not_stops(draft):
    version, programme = draft
    _entry(version, programme, "DUP-100", "DEMO One")
    _entry(version, programme, "DUP-100", "DEMO Two")
    _entry(version, programme, "DUP-101", "DEMO One")
    report = svc.validate_version(version)
    assert report.ok, _messages(report)
    assert any("DUP-100 is printed more than once" in w["message"] for w in report.warnings)


def test_validation_detects_missing_programme_and_bad_placement(draft):
    version, programme = draft
    _entry(version, None, "ORPH-100")
    _entry(version, programme, "SEM-100", semester=7)
    _entry(version, programme, "LVL-100", level=7)
    messages = _messages(svc.validate_version(version))
    assert "department or programme is missing" in messages
    assert "semester 7 is invalid" in messages
    assert "NTA level 7 is not defined" in messages


def test_validation_detects_missing_fields_and_broken_prerequisites(draft):
    version, programme = draft
    _entry(version, programme, None, name=None)
    _entry(version, programme, "PRE-100", prerequisite_codes="NOPE-999")
    _entry(version, programme, "CYC-1", "DEMO Cycle A", prerequisite_codes="CYC-2")
    _entry(version, programme, "CYC-2", "DEMO Cycle B", prerequisite_codes="CYC-1")
    _entry(version, programme, "SELF-1", "DEMO Self", prerequisite_codes="SELF-1")
    messages = _messages(svc.validate_version(version))
    assert "module code is required" in messages and "module name is required" in messages
    assert "NOPE-999 is not in this prospectus" in messages
    assert "Prerequisites form a loop" in messages
    assert "cannot be its own prerequisite" in messages


def test_valid_version_passes_with_credit_warning_only(draft):
    version, programme = draft
    _entry(version, programme, "OK-100", "DEMO Base")
    _entry(version, programme, "OK-200", "DEMO Next", semester=2, prerequisite_codes="OK-100", credits=None)
    report = svc.validate_version(version)
    assert report.ok, _messages(report)
    assert any("credits not recorded" in w["message"] for w in report.warnings)


def test_docx_tables_keep_their_rows_and_place(tmp_path):
    import docx

    from file_processing import extract_text

    document = docx.Document()
    document.add_paragraph("Semester I")
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("DMO 04111", "DEMO Drawing", "6")):
        cell.text = text
    document.add_paragraph("Semester II")
    path = tmp_path / "demo.docx"
    document.save(path)
    assert extract_text(str(path), "docx") == "Semester I\nDMO 04111 DEMO Drawing 6\nSemester II"


def test_bundled_prospectus_loads_once(app, admin, tmp_path, live_restored):
    (tmp_path / "dit_prospectus_2031_2032.txt").write_text(DIT_LAYOUT, encoding="utf-8")
    with app.app_context():
        version = svc.load_bundled_prospectus(str(tmp_path))
        db.session.commit()
        assert version.status == "published" and version.academic_year_label == "2031/2032"
        assert version.label.startswith("DIT Prospectus · 2031/2032")
        version_id = version.id
        assert svc.load_bundled_prospectus(str(tmp_path)) is None
        svc.undo_last_publish(user=None)
        db.session.commit()
        assert svc.load_bundled_prospectus(str(tmp_path)) is None, "an undone prospectus is never reloaded"
    page = admin.get(f"/admin/prospectus/{version_id}")
    assert b"Notes from the read" in page.data and b"Programmes not read" in page.data
    assert b"NTA 7 \xc2\xb7 Higher Diploma" in page.data


def test_shipped_prospectus_is_the_one_the_reader_was_tuned_on():
    import os
    names = os.listdir(svc.BUNDLED_PROSPECTUS_DIR)
    assert "dit_prospectus_2025_2026.txt" in names
    with open(os.path.join(svc.BUNDLED_PROSPECTUS_DIR, "dit_prospectus_2025_2026.txt"), encoding="utf-8") as handle:
        rows, notes = svc.extract_prospectus(handle.read())
    assert len(rows) > 1400
    assert {row["department"] for row in rows} >= {"Civil Engineering", "Computer Studies", "Mwanza Campus"}
    assert all(row["nta_level"] and row["semester"] for row in rows)


def test_students_and_lecturers_get_no_dead_links_after_modules_retire(app, admin, student, lecturer, seeded_paths,
                                                                       live_restored):
    p = seeded_paths
    with app.app_context():
        retired = [r.id for r in Resource.query.filter(Resource.module_id != p["module_id"],
                                                       Resource.verification_status == "verified")]
    student.get(f"/resource/{retired[0]}")  # history that points at a module about to retire
    _upload(admin, _csv([(p["department"], p["programme"], p["level"], p["semester"], p["code"],
                        "DEMO Kept Module", 10, "")]))
    page = student.get("/dashboard").get_data(as_text=True)
    assert not any(f'href="/resource/{rid}"' in page for rid in retired)
    with app.app_context():
        department = User.query.filter_by(email="student@dit.ac.tz").one().department
    programmes = student.get(f"/archive/{department.slug}").get_data(as_text=True)
    assert programmes.count('<a class="entity-card') == 1
    assert "/resource/" not in lecturer.get("/lecturer/resources").get_data(as_text=True)
